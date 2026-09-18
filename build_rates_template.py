#!/usr/bin/env python3
"""Build an empty rates template Excel from the quotation portal structure."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / ".xlsxlib"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "Ellipse_Solar_Rates_Template.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="E8621F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="FFF3EB")
HINT_FONT = Font(italic=True, color="666666", size=10)
THIN = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
PRICE_FILL = PatternFill("solid", fgColor="FFFDE7")

PANEL_BRANDS = [
    "INA",
    "Adani",
    "Waaree",
    "Tata",
    "Renewsys",
    "Premier",
    "Vikram",
    "Goldi",
    "Saatvik",
    "Jakson",
    "Canadian Solar",
    "Jinko",
    "Trina",
]
TECHS = ["MonoPERC", "TOPCon"]
DCR = ["DCR", "Non-DCR"]
KW_SIZES = list(range(3, 21))

ONGIRD_TIERS = [
    (2, "1Ph", 1),
    (3, "1Ph", 1),
    (3.6, "1Ph", 1),
    (4, "1Ph", 1),
    (4.6, "1Ph", 1),
    (5, "1Ph", 1),
    (5.5, "1Ph", 1),
    (6, "1Ph", 1),
    (6, "1Ph", 2),
    (5, "3Ph", 1),
    (6, "3Ph", 1),
    (8, "3Ph", 1),
    (10, "3Ph", 1),
    (10, "3Ph", 2),
    (12, "3Ph", 2),
    (15, "3Ph", 2),
    (20, "3Ph", 2),
    (25, "3Ph", 2),
    (30, "3Ph", 2),
    (40, "3Ph", 3),
    (50, "3Ph", 4),
    (60, "3Ph", 4),
    (75, "3Ph", 4),
    (100, "3Ph", 8),
    (110, "3Ph", 9),
    (125, "3Ph", 9),
]
HYBRID_OFFGRID_TIERS = [
    (3, "1Ph", 1),
    (3.6, "1Ph", 1),
    (5, "1Ph", 2),
    (6, "1Ph", 2),
    (5, "3Ph", 2),
    (6, "3Ph", 2),
    (8, "3Ph", 2),
    (10, "3Ph", 2),
    (12, "3Ph", 2),
    (15, "3Ph", 2),
    (20, "3Ph", 3),
]
INVERTER_BRANDS = {
    "On-Grid": ["Polycab", "Growatt", "Sungrow", "Deye", "Solis"],
    "Hybrid": ["Polycab", "Deye", "Growatt", "Luminous", "Eastman"],
    "Off-Grid": ["Polycab", "Luminous", "Microtek", "Sukam", "Eastman"],
}

BATTERY_BRANDS = ["Exide", "Luminous", "Livguard", "Amaron", "Okaya", "Eastman", "SF Sonic"]
BATTERY_SIZES = [(100, 12), (120, 12), (150, 12), (180, 12), (200, 12), (220, 12), (100, 24), (150, 24), (200, 24)]

BI_BRANDS = ["Luminous", "Exide", "Polycab", "Microtek", "Sukam", "Eastman"]
BI_KW = [1, 2, 3, 5, 7.5, 10]


def write_sheet(ws, headers, rows, note, price_cols=()):
    col_count = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    note_cell = ws.cell(1, 1, note)
    note_cell.font = HINT_FONT
    note_cell.fill = SECTION_FILL
    note_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48

    for col, title in enumerate(headers, 1):
        cell = ws.cell(2, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = "A3"
    last_row = 2 + max(len(rows), 1)
    ws.auto_filter.ref = f"A2:{get_column_letter(col_count)}{last_row}"

    price_idx = {headers.index(c) + 1 for c in price_cols if c in headers}
    for r, row in enumerate(rows, 3):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val if val is not None else None)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if c in price_idx:
                cell.fill = PRICE_FILL
                cell.number_format = "0.00"
    for i, title in enumerate(headers, 1):
        width = max(14, min(38, len(title) + 4))
        if rows:
            sample = max((len(str(r[i - 1])) if r[i - 1] is not None else 0) for r in rows[:40])
            width = max(width, min(38, sample + 3))
        ws.column_dimensions[get_column_letter(i)].width = width


def main():
    wb = Workbook()

    # --- Read me ---
    ws = wb.active
    ws.title = "Read_Me"
    ws["A1"] = "Ellipse Solar — Rates template"
    ws["A1"].font = Font(bold=True, size=16, color="E8621F")
    notes = [
        "",
        "How to fill this file",
        "1. Yellow cells are prices / amounts. Leave a cell blank if that item is not sold.",
        "2. Do not rename sheet names or column headers — those are used to update rates.json later.",
        "3. Extra brands and sizes are already listed. Fill only what you actually sell.",
        "4. Profit is a fixed rupee amount for that brand + system kW (not multiplied by kW).",
        "5. Panel wattages: put Y under Available for each wattage you stock. Put Y under Default for one wattage per brand + technology.",
        "6. Inverters: fill Price for each kW / phase / MPPT you sell. Blank rows are ignored.",
        "7. When done, send this file back and the portal rates.json will be updated from it.",
        "",
        "Sheets",
        "Panel_Rates — ₹/watt for each brand, MonoPERC/TOPCon, DCR/Non-DCR, plus warranty years",
        "Panel_Wattages — which module wattages exist for each brand + technology",
        "Profit — static profit amount for each brand + system kW",
        "Inverters — on-grid / hybrid / off-grid inverter prices by brand and size",
        "Batteries — battery bank unit prices",
        "Battery_Inverters — extra inverter used when on-grid + battery backup",
        "Structure — GI structure cost by system kW",
        "Components — BOS items (wires, meters, earthing, etc.)",
        "Other_Charges — GST, transport, DISCOM, civil work, subsidy, loan constants",
    ]
    for i, line in enumerate(notes, 1):
        ws.cell(i, 1, line)
        if line in ("How to fill this file", "Sheets"):
            ws.cell(i, 1).font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 110

    # --- Panel rates ---
    headers = ["Brand", "Technology", "DCR_Type", "Price_Per_Watt", "Warranty_Years"]
    rows = []
    for brand in PANEL_BRANDS:
        for tech in TECHS:
            for dcr in DCR:
                rows.append([brand, tech, dcr, None, None])
    for _ in range(8):
        rows.append([None, None, None, None, None])
    ws = wb.create_sheet("Panel_Rates")
    write_sheet(
        ws,
        headers,
        rows,
        "Fill Price_Per_Watt (₹) and Warranty_Years. Extra blank rows at the bottom are for new brands. Leave a row blank if not sold.",
        price_cols=("Price_Per_Watt", "Warranty_Years"),
    )

    # --- Panel wattages ---
    common_watts = {
        "MonoPERC": [530, 535, 540, 545, 550, 555, 560, 565],
        "TOPCon": [570, 575, 580, 585, 590, 595, 600, 605, 610, 615, 620, 625, 630, 635, 640, 650, 660, 670, 700],
    }
    headers = ["Brand", "Technology", "Wattage", "Available_Y_N", "Default_Y_N"]
    rows = []
    for brand in PANEL_BRANDS:
        for tech in TECHS:
            for w in common_watts[tech]:
                rows.append([brand, tech, w, None, None])
    ws = wb.create_sheet("Panel_Wattages")
    write_sheet(
        ws,
        headers,
        rows,
        "Put Y in Available_Y_N for wattages you stock. Put Y in Default_Y_N for exactly one wattage per Brand + Technology.",
    )

    # --- Profit ---
    headers = ["Brand", "System_kW", "Profit_Amount"]
    rows = []
    for brand in PANEL_BRANDS:
        for kw in KW_SIZES:
            rows.append([brand, kw, None])
    ws = wb.create_sheet("Profit")
    write_sheet(
        ws,
        headers,
        rows,
        "Static profit in ₹ for that brand and system kW. Example: INA 3 kW = 45000 means the quote adds ₹45,000, not 3 × 45000.",
        price_cols=("Profit_Amount",),
    )

    # --- Inverters ---
    headers = ["System_Type", "Brand", "Capacity_kW", "Phase", "MPPT", "Price", "Warranty_Years"]
    rows = []
    for system, brands in INVERTER_BRANDS.items():
        tiers = ONGIRD_TIERS if system == "On-Grid" else HYBRID_OFFGRID_TIERS
        for brand in brands:
            for kw, phase, mppt in tiers:
                rows.append([system, brand, kw, phase, mppt, None, None])
    ws = wb.create_sheet("Inverters")
    write_sheet(
        ws,
        headers,
        rows,
        "Fill Price for each inverter you sell. Hybrid and Off-Grid have their own brands and sizes. Blank price = not sold.",
        price_cols=("Price", "Warranty_Years"),
    )

    # --- Batteries ---
    headers = ["Brand", "AH", "Voltage", "Label", "Price", "Warranty_Years"]
    rows = []
    for brand in BATTERY_BRANDS:
        for ah, v in BATTERY_SIZES:
            rows.append([brand, ah, v, f"{ah} AH {v}V", None, None])
    ws = wb.create_sheet("Batteries")
    write_sheet(
        ws,
        headers,
        rows,
        "Unit price per battery. Quantity is chosen on the quotation form. Usable kWh on the portal is AH × V × 83.33%.",
        price_cols=("Price", "Warranty_Years"),
    )

    # --- Battery inverters ---
    headers = ["Brand", "Capacity_kW", "Label", "Price", "Warranty_Years"]
    rows = []
    for brand in BI_BRANDS:
        for kw in BI_KW:
            rows.append([brand, kw, f"{kw} kW", None, None])
    ws = wb.create_sheet("Battery_Inverters")
    write_sheet(
        ws,
        headers,
        rows,
        "Used only for On-Grid systems that also have battery backup. This is the extra inverter for the battery bank.",
        price_cols=("Price", "Warranty_Years"),
    )

    # --- Structure ---
    headers = ["System_kW", "Price", "Brand_Notes"]
    rows = [[kw, None, "TATA / Apollo GI"] for kw in KW_SIZES]
    ws = wb.create_sheet("Structure")
    write_sheet(
        ws,
        headers,
        rows,
        "Complete GI structure cost for that system size. Currently one price table is shared by TATA and Apollo.",
        price_cols=("Price",),
    )

    # --- Components ---
    headers = ["Item", "Variant", "Unit", "Price", "Default_Qty", "Notes"]
    component_rows = [
        ["ACDB + DCDB Combo", "1 Phase 1-6 kW", "set", None, 1, "Polycab"],
        ["ACDB + DCDB Combo", "3 Phase 5-10 kW", "set", None, 1, "Polycab"],
        ["ACDB + DCDB Combo", "3 Phase 10-20 kW", "set", None, 1, "Polycab"],
        ["Chemical Earthing 1 Mtr", "Standard", "nos", None, 3, ""],
        ["Lightning Arrester", "Standard", "nos", None, 1, ""],
        ["Earthing Wire", "1 Phase 2.5mm", "₹ per meter", None, None, "Base 50m + 15m per extra floor"],
        ["Earthing Wire", "3 Phase 4mm", "₹ per meter", None, None, "Base 50m + 15m per extra floor"],
        ["DC Wire 4mm", "Copper", "₹ per meter", None, None, "Base 30m + 10m per extra floor"],
        ["Black Pipe GI", "Standard", "₹ per meter", None, None, "Same length as earthing wire"],
        ["Chemical Bags", "Earthing", "nos", None, 1, ""],
        ["AC Armoured Cable", "1 Phase 2 core 6 sq mm", "₹ per meter", None, 5, "Fixed 5m"],
        ["AC Armoured Cable", "3 Phase 4 core 10 sq mm", "₹ per meter", None, 5, "Fixed 5m"],
        ["PVC Ducting 45x45 1Mtr", "Standard", "nos", None, 1, ""],
        ["MC4 Connectors", "Pair", "₹ per pair", None, None, "Quantity = panel count + 4"],
        ["Solar Meter", "1 Phase", "nos", None, 1, "L&T / HPL"],
        ["Solar Meter", "3 Phase", "nos", None, 1, "L&T / HPL"],
        ["Net Meter", "1 Phase", "nos", None, 1, "Currently not charged; fill if you want it back"],
        ["Net Meter", "3 Phase", "nos", None, 1, "Currently not charged; fill if you want it back"],
    ]
    ws = wb.create_sheet("Components")
    write_sheet(
        ws,
        headers,
        component_rows,
        "Balance of system. Fill Price. Default_Qty is used when quantity is not calculated from floors or panel count.",
        price_cols=("Price", "Default_Qty"),
    )

    # --- Other charges ---
    headers = ["Category", "Item", "Value", "Notes"]
    other = [
        ["Tax", "GST rate (decimal)", None, "Example: 0.089 for 8.9%"],
        ["Additional", "Vendor login charges", None, "₹"],
        ["Additional", "Transportation", None, "₹"],
        ["Additional", "Miscellaneous", None, "₹"],
        ["Additional", "DISCOM", None, "₹"],
        ["Additional", "Civil work material per kW", None, "₹ per kW × system size"],
        ["Subsidy", "Center subsidy 1 kW", None, "₹"],
        ["Subsidy", "Center subsidy 2 kW", None, "₹"],
        ["Subsidy", "Center subsidy 3 kW and above", None, "₹"],
        ["Subsidy", "State subsidy", None, "₹"],
        ["Battery", "Usable kWh factor", None, "Example: 0.8333 for 83.33% of AH×V"],
        ["Loan", "Max loan amount", None, "₹"],
        ["Loan", "Interest rate % annual", None, "Example: 5.75"],
        ["Loan", "Default tenure years", None, ""],
        ["Loan", "Units generated per kW per day", None, "Used in bill recommendation and loan savings"],
        ["Loan", "Average electricity rate ₹/unit", None, ""],
        ["Loan", "Govt buyback rate ₹/unit", None, ""],
        ["Structure brand", "TATA", None, "Y if offered"],
        ["Structure brand", "Apollo", None, "Y if offered"],
        ["Structure brand", "Jindal", None, "Y if offered"],
        ["Sales category", "S1 to S20 step amount", None, "Currently ₹1000 per step; fill to change"],
        ["Discount", "D1 to D20 step amount", None, "Currently ₹1000 per step; fill to change"],
    ]
    ws = wb.create_sheet("Other_Charges")
    write_sheet(
        ws,
        headers,
        other,
        "Portal-wide numbers. Yellow Value column is what gets copied into rates.json / loan settings.",
        price_cols=("Value",),
    )

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
